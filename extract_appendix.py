import sys
try:
    import docx
except ImportError:
    print("python-docx not installed.")
    sys.exit(1)

doc = docx.Document('Project_Report_Expanded.docx')
in_appendix = False
text = []
for p in doc.paragraphs:
    if 'Appendix A' in p.text or 'APPENDIX A' in p.text:
        in_appendix = True
    if in_appendix:
        text.append(p.text)

with open('appendix_a.txt', 'w', encoding='utf-8') as f:
    f.write(f"Extracted {len(text)} paragraphs for Appendix A.\n")
    if len(text) > 0:
        for t in text:
            if t.strip():
                f.write(t + "\n")
    else:
        f.write("Appendix A not found in paragraphs. Checking tables...\n")
        cnt = 0
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if 'Appendix A' in cell.text or 'APPENDIX A' in cell.text:
                        f.write("Found in table:\n")
                        f.write(cell.text + "\n")
                        cnt += 1
                        if cnt > 5:
                            break
