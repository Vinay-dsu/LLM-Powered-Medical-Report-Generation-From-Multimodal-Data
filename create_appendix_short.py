from docx import Document
from docx.shared import Pt
import os

doc = Document()

heading = doc.add_heading("Appendix A: Core Implementation Source Code", level=1)
doc.add_paragraph("This appendix contains the core source code for the Medical Report Generator (V5), encompassing the multimodal model architecture (ViT + BioGPT) and the Flask web service.")

def add_code_file(doc, title, filepath):
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    with open(filepath, 'r', encoding='utf-8') as f:
        # We can extract just the critical sections from app.py
        # Or just write the whole thing - let's write the whole app.py, skipping the long string formatting function parse_report_sections 
        lines = f.readlines()
        filtered_lines = []
        skip = False
        for line in lines:
            if line.startswith("def parse_report_sections(raw_text):"):
                skip = True
                filtered_lines.append(line)
                filtered_lines.append('    # ... omitted string formatting logic for brevity ...\n')
                filtered_lines.append('    return {"findings": [], "impression": [], "recommendations": []}\n\n')
                continue
            
            if skip:
                if line.startswith("def ") or line.startswith("# ========================"):
                    skip = False
                else:
                    continue
            
            filtered_lines.append(line)
            
        code_text = "".join(filtered_lines)
    
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)

add_code_file(doc, "app.py (Flask Backend & Model Definition Worker)", r"c:\GoatProject\app.py")

doc.save(r"c:\GoatProject\Appendix_A_Shortened.docx")
doc.save(r"c:\GoatProject\Appendix_A.docx") # Overwrite
print("Successfully generated Shortened Appendix_A.docx")
