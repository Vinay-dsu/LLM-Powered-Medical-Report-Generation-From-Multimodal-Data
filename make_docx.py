import docx
import re
from docx.shared import Pt

def add_markdown_paragraph(doc, text, style=None):
    if style:
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph()
        
    # Super basic markdown bold parser
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)

doc = docx.Document()

with open(r"C:\Users\vinay\.gemini\antigravity\brain\70a91ec1-ba2d-42e9-8730-9c0f62b696ae\expanded_report.md", "r", encoding="utf-8") as f:
    content = f.read()

# Split by double newline to get distinct paragraphs/elements
blocks = content.split("\n\n")

for block in blocks:
    block = block.strip()
    if not block:
        continue
    
    # Check if lines have lists
    lines = block.split('\n')
    if len(lines) > 1 and all(l.startswith('- ') or l.startswith('* ') for l in lines):
        for l in lines:
            add_markdown_paragraph(doc, l[2:], style='List Bullet')
        continue
    
    if len(lines) > 1 and all(re.match(r'^\d+\.\s', l) for l in lines):
        for l in lines:
            add_markdown_paragraph(doc, re.sub(r'^\d+\.\s', '', l), style='List Number')
        continue

    # Normal parsing
    if block.startswith("# "):
        doc.add_heading(block[2:], level=1)
    elif block.startswith("## "):
        doc.add_heading(block[3:], level=2)
    elif block.startswith("### "):
        doc.add_heading(block[4:], level=3)
    elif block.startswith("- ") or block.startswith("* "):
         add_markdown_paragraph(doc, block[2:], style='List Bullet')
    elif re.match(r'^\d+\.\s', block):
         add_markdown_paragraph(doc, re.sub(r'^\d+\.\s', '', block), style='List Number')
    else:
         add_markdown_paragraph(doc, block)

doc.save("Project_Report_Expanded_Final.docx")
print("Saved to Project_Report_Expanded_Final.docx")
