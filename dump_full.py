import sys
try:
    import docx
except ImportError:
    print("python-docx not installed.")
    sys.exit(1)

doc = docx.Document('Project_Report_Expanded.docx')

with open('full_report.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"{i}: {p.text}\n")
    f.write("---TABLES---\n")
    for i, table in enumerate(doc.tables):
        f.write(f"Table {i}:\n")
        for row in table.rows:
            for cell in row.cells:
                f.write(cell.text + "\n")
