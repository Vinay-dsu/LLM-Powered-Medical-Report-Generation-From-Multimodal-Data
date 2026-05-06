from docx import Document
from docx.shared import Pt
import os

doc = Document()

# Add Title
heading = doc.add_heading("Appendix A: Core Implementation Source Code", level=1)
doc.add_paragraph("This appendix contains the core source code for the proposed multimodal medical imaging architecture, including the Flask application backend and the generative inference engine.")

def add_code_file(doc, filename, filepath):
    if not os.path.exists(filepath):
        print(f"File {filepath} not found.")
        return
    
    doc.add_heading(filename, level=2)
    p = doc.add_paragraph()
    with open(filepath, 'r', encoding='utf-8') as f:
        code_text = f.read()
    
    # We add run and format as 'Courier New' code style
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)

files_to_include = [
    ("app.py (Flask Backend Application)", "app.py"),
    ("generate_report.py (Inference Engine & Model Architecture)", "generate_report.py"),
    ("run_local.py (Local Evaluation Script)", "run_local.py"),
    ("evaluate_metrics.py (NLP Evaluation and Metrics Script)", "evaluate_metrics.py")
]

for title, filename in files_to_include:
    print(f"Adding {filename} to Appendix_A...")
    add_code_file(doc, title, os.path.join(r"c:\GoatProject", filename))

doc.save(r"c:\GoatProject\Appendix_A.docx")
print("Successfully generated Appendix_A.docx")
